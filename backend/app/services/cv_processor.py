import io
import logging
from typing import Optional

import PyPDF2
from docx import Document

logger = logging.getLogger(__name__)

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class CVProcessor:
    """Handles text extraction from CV files (PDF and DOCX)."""

    def extract_text_from_pdf(self, pdf_bytes: bytes) -> str:
        """Extract plain text from a PDF file given as raw bytes."""
        try:
            reader = PyPDF2.PdfReader(io.BytesIO(pdf_bytes))
            pages_text = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    pages_text.append(text)
            return "\n".join(pages_text)
        except Exception as exc:
            logger.error("Failed to extract text from PDF: %s", exc)
            raise ValueError(f"PDF extraction failed: {exc}") from exc

    def extract_text_from_docx(self, docx_bytes: bytes) -> str:
        """Extract plain text from a Word (.docx) document given as raw bytes."""
        try:
            doc = Document(io.BytesIO(docx_bytes))
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text:
                            paragraphs.append(cell_text)
            return "\n".join(paragraphs)
        except Exception as exc:
            logger.error("Failed to extract text from DOCX: %s", exc)
            raise ValueError(f"DOCX extraction failed: {exc}") from exc

    def extract_text(self, file_bytes: bytes, mime_type: str) -> str:
        """Route to the correct extractor based on MIME type."""
        if mime_type == PDF_MIME or mime_type == "pdf":
            return self.extract_text_from_pdf(file_bytes)
        elif mime_type == DOCX_MIME or mime_type in ("docx", "word"):
            return self.extract_text_from_docx(file_bytes)
        else:
            # Try PDF first, then DOCX
            try:
                return self.extract_text_from_pdf(file_bytes)
            except Exception:
                pass
            try:
                return self.extract_text_from_docx(file_bytes)
            except Exception:
                pass
            raise ValueError(f"Unsupported MIME type and extraction failed: {mime_type}")

    def truncate_text(self, text: str, max_chars: int = 12000) -> str:
        """Truncate CV text to avoid exceeding AI model token limits."""
        if len(text) <= max_chars:
            return text
        return text[:max_chars] + "\n[... text truncated for processing ...]"
